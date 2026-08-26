"""Excel export helpers for attendance and shifts."""
from io import BytesIO

import openpyxl


def build_attendance_excel(records: list) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(["Date", "Day", "Employee", "Designation", "Project", "Shift Code", "Shift Name",
               "Work Type", "Shift Status", "Check In", "Check Out", "Total Hours", "Late",
               "Office", "Manual Entry", "Remarks"])
    for r in records:
        ws.append([
            r.get("attendanceDate") or "",
            r.get("dayName") or "",
            r.get("userName") or "",
            r.get("userDesignation") or "",
            r.get("shiftProject") or "",
            r.get("shiftCode") or "",
            r.get("shiftName") or "",
            r.get("workLocationStatus") or "",
            r.get("derivedStatus") or r.get("status") or "",
            r.get("checkInTime") or "",
            r.get("checkOutTime") or "",
            r.get("totalHours") if r.get("totalHours") is not None else "",
            "Yes" if r.get("isLate") else "No",
            r.get("officeName") or "",
            "Yes" if r.get("isManualEntry") else "No",
            r.get("remarks") or "",
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def build_employees_excel(records: list) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(["First Name", "Middle Name", "Last Name", "Email", "Employee ID",
               "Designation", "Status", "Onboarded", "Work Location", "Country",
               "Contact", "Alt Contact", "Reports To", "Joining Date",
               "Projects", "Roles", "Managers"])
    for r in records:
        projects = ", ".join(r.get("projects") or [])
        roles = ", ".join(
            f"{x['project']}: {x['role']}" for x in (r.get("redmine_roles") or [])
        )
        managers = ", ".join(sorted({
            m for x in (r.get("redmine_roles") or []) for m in (x.get("managers") or [])
        }))
        ws.append([
            r.get("first_name") or "",
            r.get("middle_name") or "",
            r.get("last_name") or "",
            r.get("user_email") or "",
            r.get("employee_id") or "",
            r.get("designation") or "",
            r.get("status") or "",
            "Yes" if r.get("onboarded_at") else "No",
            r.get("work_location_status") or "",
            r.get("country") or "",
            r.get("contact_number") or "",
            r.get("alt_contact_number") or "",
            r.get("reports_to_name") or "",
            r.get("joining_date") or "",
            projects,
            roles,
            managers,
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def build_shifts_excel(records: list) -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shifts"
    ws.append(["Date", "Day", "Employee", "Designation", "Project", "Shift Code", "Shift Name",
               "Work Status", "Status", "Shift Start", "Shift End", "Check In", "Check Out",
               "Total Hours", "Late (hrs)", "Work Address"])
    for r in records:
        ws.append([
            r.get("date") or "",
            r.get("dayName") or "",
            r.get("userName") or "",
            r.get("userDesignation") or "",
            r.get("projectName") or "",
            r.get("shift") or "",
            r.get("shiftName") or "",
            r.get("workStatus") or "",
            r.get("status") or "",
            r.get("shiftStartTime") or "",
            r.get("shiftEndTime") or "",
            r.get("checkInTime") or "",
            r.get("checkOutTime") or "",
            r.get("totalHours") if r.get("totalHours") is not None else "",
            r.get("lateByHours") if r.get("lateByHours") is not None else "",
            r.get("workAddress") or "",
        ])
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
